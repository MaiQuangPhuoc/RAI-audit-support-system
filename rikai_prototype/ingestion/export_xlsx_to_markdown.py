"""
export_xlsx_to_markdown.py
Script trích xuất nội dung từ file Excel (.xlsx) đa sheet, xử lý bảng phức tạp (merged cells),
và lưu kết quả thành file Markdown (.md).

Nguyên tắc:
- Chạy hoàn toàn bằng code Python thuần (openpyxl), KHÔNG gọi LLM.
- Giữ nguyên cấu trúc phân cấp bằng cách unmerge và điền giá trị lan truyền cho ô gộp.
- Giữ nguyên ô trống (để phục vụ Hearing Sheet / kiểm toán).
"""

import sys
import os
from typing import List, Dict, Any
import openpyxl


def unmerge_and_propagate(sheet) -> None:
    """
    Xử lý ô gộp (merged cells): sao chép giá trị ô đầu cho toàn bộ vùng merge
    để giữ nguyên ngữ cảnh danh mục, số thứ tự và tiêu chí kiểm toán.
    """
    merged_ranges = list(sheet.merged_cells.ranges)
    for rng in merged_ranges:
        min_col, min_row, max_col, max_row = rng.min_col, rng.min_row, rng.max_col, rng.max_row
        top_left_value = sheet.cell(row=min_row, column=min_col).value
        sheet.unmerge_cells(str(rng))
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                sheet.cell(row=row, column=col).value = top_left_value


def escape_md_cell(val: Any) -> str:
    """Xử lý ký tự đặc biệt trong bảng Markdown."""
    if val is None:
        return ""
    text = str(val).strip()
    # Thay thế dấu xuống dòng và ký tự pipe |
    text = text.replace("\r\n", " ").replace("\n", " ").replace("|", "\\|")
    return text


def convert_sheet_to_markdown(sheet_name: str, sheet) -> str:
    """Chuyển đổi một sheet trong openpyxl thành đoạn text Markdown."""
    unmerge_and_propagate(sheet)
    rows_raw = list(sheet.iter_rows(values_only=True))

    if not rows_raw:
        return f"## Sheet: {sheet_name}\n\n*(Sheet trống)*\n\n"

    # Tìm dòng bắt đầu có dữ liệu
    start_idx = 0
    while start_idx < len(rows_raw) and all(c is None for c in rows_raw[start_idx]):
        start_idx += 1

    if start_idx >= len(rows_raw):
        return f"## Sheet: {sheet_name}\n\n*(Sheet trống)*\n\n"

    # Kiểm tra số lượng cột có nội dung
    max_non_empty_cols = max(len([c for c in r if c is not None]) for r in rows_raw[start_idx:])

    md_lines: List[str] = [f"## Sheet: {sheet_name}\n"]

    # Nếu chỉ có 1-2 cột: Định dạng thành danh sách Key - Value / Text mô tả
    if max_non_empty_cols <= 2:
        for r in rows_raw[start_idx:]:
            non_empty = [escape_md_cell(c) for c in r if c is not None and str(c).strip()]
            if len(non_empty) == 1:
                md_lines.append(f"- **{non_empty[0]}**")
            elif len(non_empty) > 1:
                md_lines.append(f"- **{non_empty[0]}**: {non_empty[1]}")
        md_lines.append("\n")
        return "\n".join(md_lines)

    # Nếu là bảng biểu kiểm toán (nhiều cột)
    header_raw = rows_raw[start_idx]
    headers = [escape_md_cell(c) if c is not None and str(c).strip() else f"Cột {i+1}" for i, c in enumerate(header_raw)]
    
    # Tạo header bảng Markdown
    md_lines.append("| " + " | ".join(headers) + " |")
    md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")

    # Điền các dòng dữ liệu
    for row in rows_raw[start_idx + 1:]:
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        row_cells = []
        for i in range(len(headers)):
            val = row[i] if i < len(row) else ""
            row_cells.append(escape_md_cell(val))
        md_lines.append("| " + " | ".join(row_cells) + " |")

    md_lines.append("\n")
    return "\n".join(md_lines)


def export_xlsx_to_md(xlsx_path: str, output_md_path: str = None) -> str:
    """Hàm chính đọc file XLSX và ghi ra file Markdown."""
    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(f"Không tìm thấy file tại đường dẫn: {xlsx_path}")

    if output_md_path is None:
        base, _ = os.path.splitext(xlsx_path)
        output_md_path = f"{base}_extracted.md"

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    md_content = [
        f"# Nội Dung Trích Xuất File: `{os.path.basename(xlsx_path)}`",
        f"- **Đường dẫn nguồn**: `{xlsx_path}`",
        f"- **Tổng số sheet**: {len(wb.sheetnames)}",
        f"- **Trạng thái trích xuất**: Dữ liệu thô (Không qua LLM, đã chuẩn hoá merged cells)\n",
        "---\n"
    ]

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        sheet_md = convert_sheet_to_markdown(sheet_name, sheet)
        md_content.append(sheet_md)

    final_text = "\n".join(md_content)

    with open(output_md_path, "w", encoding="utf-8") as f:
        f.write(final_text)

    return output_md_path


if __name__ == "__main__":
    file_input = (
        r"D:\PHUOC\RAI_LLM\rikai_prototype2\rikai_prototype\data\input\xlsx_data_llm.xlsx"
    )

    file_output = (
        r"D:\PHUOC\RAI_LLM\rikai_prototype2\rikai_prototype\data\input\xlsx_data_llm.md"
    )


    print(f"=== Đang xử lý trích xuất sang Markdown: {file_input} ===")
    try:
        out_path = export_xlsx_to_md(file_input, file_output)
        print(f"✓ Trích xuất thành công!")
        print(f"✓ File Markdown đã được tạo tại: {out_path}")
    except Exception as err:
        print(f"✗ Lỗi khi trích xuất file: {err}")