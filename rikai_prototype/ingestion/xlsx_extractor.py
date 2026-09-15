"""Bóc tách nội dung thô từ file XLSX — dump từng dòng thành text để LLM đọc,
không tự diễn giải cấu trúc ở bước này."""
import openpyxl


def extract_text(file_path: str) -> str:
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    lines = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        lines.append(f"[Sheet: {sheet_name}]")
        for row in sheet.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            cells = [str(c).strip() for c in row if c is not None]
            lines.append(" | ".join(cells))
    return "\n".join(lines)
