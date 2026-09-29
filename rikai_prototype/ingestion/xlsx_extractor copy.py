"""
Bóc tách nội dung thô từ file XLSX — PHÂN BIỆT được dòng text (tiêu đề/mô tả,
chỉ có 1 ô có giá trị) và dòng thuộc bảng (>= 2 ô có giá trị), để tránh dòng
text bị "ăn" theo độ rộng cột của bảng bên dưới (gây nhiễu, tốn token).

Không tự diễn giải Ý NGHĨA nội dung ở bước này — chỉ tổ chức lại cho đúng
hình dạng (text thuần vs bảng), dữ liệu vẫn nguyên văn.
"""
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def _build_merge_value_map(sheet) -> dict[tuple[int, int], str]:
    """
    openpyxl không cho ghi trực tiếp vào các ô nằm trong vùng merge (trừ ô
    trên-cùng-trái) — MergedCell.value là read-only. Nên thay vì ghi đè, xây
    1 map (dòng, cột) -> giá trị của ô gốc, để dùng khi đọc dữ liệu.
    """
    merge_map: dict[tuple[int, int], str] = {}
    for merged_range in sheet.merged_cells.ranges:
        top_value = sheet.cell(row=merged_range.min_row, column=merged_range.min_col).value
        if top_value is None:
            continue
        for r in range(merged_range.min_row, merged_range.max_row + 1):
            for c in range(merged_range.min_col, merged_range.max_col + 1):
                merge_map[(r, c)] = top_value
    return merge_map


def _read_rows_with_position(sheet, merge_map: dict[tuple[int, int], str]) -> list[list[tuple[int, str]]]:
    """Đọc từng dòng, chỉ giữ lại các ô CÓ giá trị (kể cả giá trị suy ra từ
    vùng merge), kèm số cột gốc của ô đó."""
    rows = []
    for row in sheet.iter_rows():
        cells = []
        for cell in row:
            val = cell.value
            if val is None:
                val = merge_map.get((cell.row, cell.column))
            if val is not None and str(val).strip() != "":
                cells.append((cell.column, str(val).strip()))
        rows.append(cells)
    return rows


def _group_rows_into_blocks(rows: list[list[tuple[int, str]]]) -> list[list]:
    """
    Gom các dòng thành block:
    - Dòng rỗng hoàn toàn: kết thúc block bảng đang gom (nếu có), không tạo block.
    - Dòng chỉ có 1 ô: 1 block text riêng (đứng độc lập).
    - Dòng có >= 2 ô: gom vào block bảng đang mở (liên tiếp nhau).
    """
    blocks = []
    current_table: list = []

    def flush_table():
        if current_table:
            blocks.append(("table", current_table.copy()))
            current_table.clear()

    for cells in rows:
        if not cells:
            flush_table()
            continue

        distinct_values = {val for _, val in cells}
        if len(cells) == 1 or len(distinct_values) == 1:
            # 1 ô có giá trị, HOẶC nhiều ô nhưng cùng 1 giá trị (thường do merge
            # nhiều cột làm tiêu đề) -> coi là dòng text, không phải dòng bảng.
            flush_table()
            blocks.append(("text", [[cells[0]]]))
        else:
            current_table.append(cells)
    flush_table()
    return blocks


def _render_text_block(block) -> str:
    return block[0][0][1]


def _render_table_block(block) -> str:
    columns = sorted({col for row in block for col, _ in row})
    if not columns:
        return ""

    header_map = dict(block[0])
    data_rows = block[1:]

    def col_label(c):
        # Ưu tiên dùng chính nội dung dòng đầu làm tên cột (thường là header thật
        # của bảng). Nếu ô đó trống, dùng ký hiệu cột Excel (A, B, C...) làm nhãn.
        return header_map.get(c) or get_column_letter(c)

    lines = [
        "| " + " | ".join(col_label(c) for c in columns) + " |",
        "| " + " | ".join(["---"] * len(columns)) + " |",
    ]
    for row in data_rows:
        row_map = dict(row)
        lines.append("| " + " | ".join(row_map.get(c, "") for c in columns) + " |")
    return "\n".join(lines)


def extract_text(file_path: str) -> str:
    workbook = load_workbook(file_path, data_only=True)
    output_parts = []

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        merge_map = _build_merge_value_map(sheet)

        output_parts.append(f"## Sheet: {sheet_name}")

        rows = _read_rows_with_position(sheet, merge_map)
        blocks = _group_rows_into_blocks(rows)

        for block_type, block in blocks:
            if block_type == "text":
                output_parts.append(_render_text_block(block))
            else:
                rendered = _render_table_block(block)
                if rendered:
                    output_parts.append(rendered)

    return "\n\n".join(output_parts)


def extract_structured(file_path: str) -> list:
    """
    Bóc tách trực tiếp thành list[SheetTable] — KHÔNG qua LLM — để đảm bảo giữ
    ĐÚNG 100% cấu trúc cột và dữ liệu gốc khi Auditor chưa yêu cầu thay đổi gì.
    LLM chỉ nên dùng để tóm tắt/hiểu nội dung hoặc khi Auditor thật sự yêu cầu
    sửa — không dùng để tái tạo lại bảng, vì mọi bước qua LLM đều có rủi ro
    (dù nhỏ) làm lệch dữ liệu so với bản gốc.

    Logic: dùng lại đúng bộ tách block text/bảng ở trên. Mỗi sheet lấy block
    BẢNG ĐẦU TIÊN làm dữ liệu chính (columns lấy từ dòng đầu của block, đúng
    tên gốc); mọi block TEXT (mô tả, tiêu đề, tiêu chí đánh giá...) được gộp
    thành `notes` của sheet đó. Nếu 1 sheet có nhiều hơn 1 block bảng (hiếm),
    các bảng sau được giữ lại dưới dạng text trong `notes` để không mất dữ
    liệu, thay vì âm thầm bỏ qua.
    """
    from core.schemas import SheetTable  # import cục bộ để tránh vòng lặp import

    workbook = load_workbook(file_path, data_only=True)
    tables = []

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        merge_map = _build_merge_value_map(sheet)
        rows = _read_rows_with_position(sheet, merge_map)
        blocks = _group_rows_into_blocks(rows)

        notes_parts: list[str] = []
        columns: list[str] = []
        data_rows: list[dict[str, str]] = []
        table_taken = False

        for block_type, block in blocks:
            if block_type == "text":
                notes_parts.append(_render_text_block(block))
            else:
                if not table_taken:
                    col_positions = sorted({col for row in block for col, _ in row})
                    header_map = dict(block[0])
                    columns = [header_map.get(c) or get_column_letter(c) for c in col_positions]
                    for row in block[1:]:
                        row_map = dict(row)
                        data_rows.append({
                            columns[i]: row_map.get(pos, "") for i, pos in enumerate(col_positions)
                        })
                    table_taken = True
                else:
                    # Bảng thứ 2 trở lên trong cùng 1 sheet (hiếm gặp) - giữ lại
                    # dạng text trong notes thay vì bỏ mất dữ liệu.
                    rendered = _render_table_block(block)
                    if rendered:
                        notes_parts.append(rendered)

        tables.append(SheetTable(
            sheet_name=sheet_name,
            notes="\n".join(notes_parts),
            columns=columns,
            rows=data_rows,
        ))

    return tables

import os
if __name__ == "__main__":


    input_path = os.path.join(os.path.dirname(__file__), "..", "data", "input", "xlsx_data_llm.xlsx")
    output_dir = os.path.join(os.path.dirname(__file__), "..", "data", "extract")
    os.makedirs(output_dir, exist_ok=True)

    file_name = os.path.splitext(os.path.basename(input_path))[0]
    output_path = os.path.join(output_dir, f"result_{file_name}.md")

    content = extract_text(input_path)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)

    print(f"Done: {output_path}")