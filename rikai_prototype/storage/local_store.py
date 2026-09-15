"""
Lưu trữ local cho prototype:
- Lưu/đọc Hearing Sheet dạng JSON (lịch sử phiên làm việc).
- "Gửi cho Partner" = xuất Hearing Sheet ra file XLSX (2 cột: Câu hỏi/Câu trả lời)
  vào thư mục data/outbox/<partner>/.
- "Nhận từ Partner" = quét thư mục data/inbox/<partner>/ xem có file mới không.
  Đây là cách mô phỏng đơn giản cho việc gửi/nhận qua lại với Partner (thực tế
  Partner trả lời thủ công ngoài hệ thống — không thuộc phạm vi xử lý AI).

Lưu ý: việc "tự động chạy khi có file mới" theo đúng nghĩa (daemon/watcher chạy
nền) cần thư viện riêng (vd. watchdog) hoặc cron/task scheduler — ngoài phạm vi
prototype Streamlit (mô hình request-response). Ở đây dùng nút "Kiểm tra file mới"
để chủ động quét thư mục inbox mỗi khi Auditor bấm — đơn giản, đủ dùng cho demo.
"""
import glob
import json
import os
import uuid
from datetime import datetime
import os
import sys
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)
import openpyxl

import configs
from core.schemas import HearingSheet, QARow

OUTBOX_DIR = os.path.join(configs.DATA_DIR, "outbox")
INBOX_DIR = os.path.join(configs.DATA_DIR, "inbox")
SESSIONS_DIR = configs.SESSIONS_DIR

os.makedirs(OUTBOX_DIR, exist_ok=True)
os.makedirs(INBOX_DIR, exist_ok=True)


def new_session_id() -> str:
    return f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}"


def save_uploaded_file(file_bytes: bytes, filename: str, subfolder: str) -> str:
    folder = os.path.join(SESSIONS_DIR, subfolder)
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, filename)
    with open(path, "wb") as f:
        f.write(file_bytes)
    return path


# ---------------------------------------------------------------------------
# Hearing Sheet <-> JSON (lưu lịch sử) và <-> XLSX (để trao đổi với Partner)
# ---------------------------------------------------------------------------

def save_hearing_sheet_json(sheet: HearingSheet, session_id: str, version: int) -> str:
    folder = os.path.join(SESSIONS_DIR, session_id)
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"hearing_sheet_v{version}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(sheet.model_dump(), f, ensure_ascii=False, indent=2)
    return path


def export_hearing_sheet_xlsx(sheet: HearingSheet, path: str) -> str:
    """Xuất Hearing Sheet ra file XLSX 2 cột (Câu hỏi / Câu trả lời) + sheet Notes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HearingSheet"
    ws.append(["Câu hỏi", "Câu trả lời"])
    for row in sheet.rows:
        ws.append([row.question, row.answer])

    if sheet.notes.strip():
        notes_ws = wb.create_sheet("Notes")
        notes_ws["A1"] = sheet.notes

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


def send_to_partner(sheet: HearingSheet, partner: str, version: int) -> str:
    """'Gửi cho Partner' = xuất file XLSX vào data/outbox/<partner>/."""
    folder = os.path.join(OUTBOX_DIR, partner)
    filename = f"hearing_sheet_v{version}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = os.path.join(folder, filename)
    return export_hearing_sheet_xlsx(sheet, path)


def check_inbox(partner: str) -> str | None:
    """
    Quét data/inbox/<partner>/ tìm file mới nhất (xlsx/pdf). Trả về đường dẫn
    file mới nhất, hoặc None nếu chưa có gì. Trong demo, Auditor (đóng vai
    Partner) tự copy file trả lời vào đúng thư mục này.
    """
    folder = os.path.join(INBOX_DIR, partner)
    os.makedirs(folder, exist_ok=True)
    files = glob.glob(os.path.join(folder, "*.xlsx")) + glob.glob(os.path.join(folder, "*.pdf"))
    if not files:
        return None
    return max(files, key=os.path.getmtime)


def read_answered_xlsx(path: str) -> HearingSheet:
    """
    Đọc file XLSX Partner đã trả lời — CÙNG cấu trúc 2 cột đã xuất ở
    export_hearing_sheet_xlsx, nên đọc trực tiếp, không cần LLM diễn giải lại.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["HearingSheet"] if "HearingSheet" in wb.sheetnames else wb.active

    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # header
        if row and row[0]:
            question = str(row[0]).strip()
            answer = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            rows.append(QARow(question=question, answer=answer))

    notes = ""
    if "Notes" in wb.sheetnames:
        notes = str(wb["Notes"]["A1"].value or "")

    return HearingSheet(title=os.path.basename(path), rows=rows, notes=notes)


def inbox_path_for(partner: str) -> str:
    path = os.path.join(INBOX_DIR, partner)
    os.makedirs(path, exist_ok=True)
    return path
